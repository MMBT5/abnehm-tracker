import os

import pytest
from openpyxl import load_workbook

import config
import db
import excel_export


@pytest.fixture
def umgebung(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(config, "XLSX_PATH", tmp_path / "t.xlsx")
    db.init_db()
    db.speichern("2026-10-01", kcal=1730, gewicht=122.0)
    db.speichern("2026-10-08", kcal=1932, gewicht=120.4)
    return tmp_path


def test_export_inhalt(umgebung):
    assert excel_export.exportieren() is True
    wb = load_workbook(config.XLSX_PATH)
    assert wb.sheetnames == ["Tage", "Wochen"]
    tage = wb["Tage"]
    assert [c.value for c in tage[1]][:7] == ["Woche", "Tag", "Datum", "Kalorien", "Protein",
                                              "Gewicht morgens", "Bauch"]
    assert tage["B5"].value == 4 and tage["D5"].value == 1730 and tage["F5"].value == 122.0
    assert tage["D2"].value is None            # Tag 1 leer, nicht 0
    assert "A2:A8" in {str(r) for r in tage.merged_cells.ranges}
    assert wb["Wochen"]["D2"].value == 122.0


def test_gesperrte_datei_stuerzt_nicht_ab(umgebung, monkeypatch):
    excel_export.exportieren()
    vorher = config.XLSX_PATH.read_bytes()

    def gesperrt(*a, **k):
        raise PermissionError("Datei wird verwendet")
    monkeypatch.setattr(os, "replace", gesperrt)
    assert excel_export.exportieren() is False
    assert excel_export.letzter_fehler
    assert config.XLSX_PATH.read_bytes() == vorher            # Ziel unverändert
    assert [p.name for p in umgebung.iterdir() if p.suffix == ".tmp"] == []   # kein Müll

    monkeypatch.undo()
    assert excel_export.exportieren() is True                 # nächster Versuch klappt
    assert excel_export.letzter_fehler is None
